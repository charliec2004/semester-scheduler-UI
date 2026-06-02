"""Tests for workbook-first project parsing and solving."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from scheduler.cli import main
from scheduler.workbook.io import copy_template_workbook, load_workbook_project, validate_workbook_project


def _build_minimal_workbook(path: Path, *, front_desk_enabled: bool = False) -> Path:
    copy_template_workbook(path)
    workbook = load_workbook(path)

    project = workbook["Project"]
    project["B3"] = "Test Project"
    project["B4"] = front_desk_enabled

    staff = workbook["Staff"]
    staff.append(["Alice", "marketing", 6, 8, 2])

    departments = workbook["Departments"]
    departments.append(["Marketing", 6, 8])

    settings = workbook["Settings"]
    settings["C2"] = 1

    workbook.save(path)
    return path


def test_load_workbook_project_valid_minimal(tmp_path):
    workbook_path = _build_minimal_workbook(tmp_path / "minimal.xlsx")

    project = load_workbook_project(workbook_path)
    issues = validate_workbook_project(project)

    assert project.project_name == "Test Project"
    assert project.front_desk_enabled is False
    assert issues == []


def test_workbook_validation_duplicate_employees(tmp_path):
    workbook_path = _build_minimal_workbook(tmp_path / "duplicate.xlsx")
    workbook = load_workbook(workbook_path)
    workbook["Staff"].append(["Alice", "marketing", 4, 8, 3])
    workbook.save(workbook_path)

    issues = validate_workbook_project(load_workbook_project(workbook_path))

    assert any("Duplicate employee name" in issue.message for issue in issues)


def test_workbook_validation_unknown_preference_reference(tmp_path):
    workbook_path = _build_minimal_workbook(tmp_path / "unknown-pref.xlsx")
    workbook = load_workbook(workbook_path)
    preferences = workbook["Preferences"]
    preferences["A3"] = "Ghost"
    preferences["B3"] = 2.0
    workbook.save(workbook_path)

    issues = validate_workbook_project(load_workbook_project(workbook_path))

    assert any("Unknown employee 'Ghost'" in issue.message for issue in issues)


def test_workbook_validation_invalid_time_range(tmp_path):
    workbook_path = _build_minimal_workbook(tmp_path / "invalid-time.xlsx")
    workbook = load_workbook(workbook_path)
    unavailability = workbook["Unavailability"]
    unavailability.append(["Alice", "Mon", "08:00", "08:00", False, False])
    workbook.save(workbook_path)

    issues = validate_workbook_project(load_workbook_project(workbook_path))

    assert any("End time must be after start time" in issue.message for issue in issues)


def test_workbook_cli_solve_creates_output_copy(tmp_path):
    input_workbook = _build_minimal_workbook(tmp_path / "project.xlsx")
    output_workbook = tmp_path / "project-solved.xlsx"

    main(["workbook", "solve", str(input_workbook), "--output", str(output_workbook)])

    assert output_workbook.exists()

    source = load_workbook(input_workbook, data_only=True)
    solved = load_workbook(output_workbook, data_only=True)

    assert source["Results_Status"]["A3"].value == "Generated during workbook solve. Do not edit this sheet by hand."
    assert "Results_Weekly" in solved.sheetnames
    assert "Results_Formatted" in solved.sheetnames
    assert solved["Results_Status"]["B6"].value == "OPTIMAL"
    assert solved["Results_Formatted"]["A1"].value == "Title"
    assert solved["Results_Weekly"]["A3"].value == "Day"
