"""Workbook parsing, validation, and solved workbook export."""

from __future__ import annotations

import json
import math
import shutil
from collections import Counter, defaultdict
from copy import copy
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any, Iterable

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet
from ortools.sat.python import cp_model

from scheduler.config import DAY_NAMES, SLOT_MINUTES, TIME_SLOT_STARTS, hours_to_slots, is_slot_aligned_hours
from scheduler.domain.models import (
    EqualityRequest,
    FavoredEmployeeDepartment,
    ShiftTimePreference,
    TimesetRequest,
    TrainingRequest,
    normalize_department_name,
)
from scheduler.workbook.models import WorkbookProjectV1, WorkbookValidationIssue
from scheduler.workbook.settings import (
    ALL_WORKBOOK_SHEETS,
    CONSTRAINT_SECTION_HEADERS,
    FINAL_EDGE_LABEL,
    PREFERENCES_SHEET,
    PREFERENCE_SECTION_HEADERS,
    PROJECT_FIELDS,
    PROJECT_SHEET,
    README_SHEET,
    RESULT_SHEET_NAMES,
    SETTINGS_SHEET,
    SOLVER_AVAILABILITY_SHEET,
    STAFF_SHEET,
    DEPARTMENTS_SHEET,
    TIME_VALUE_OPTIONS,
    UNAVAILABILITY_SHEET,
    VALIDATION_SHEET,
    WORKBOOK_SCHEMA_VERSION,
    WORKBOOK_SETTINGS,
    WORKBOOK_SETTINGS_BY_KEY,
    WORKBOOK_TEMPLATE_FILENAME,
)

PROJECT_ROOT = Path(__file__).resolve().parents[2]
TEMPLATE_WORKBOOK_PATH = PROJECT_ROOT / WORKBOOK_TEMPLATE_FILENAME

HEADER_FILL = PatternFill(fill_type="solid", fgColor="DDEBF7")
SECTION_FILL = PatternFill(fill_type="solid", fgColor="1F4E78")
SECTION_FONT = Font(color="FFFFFF", bold=True)
HEADER_FONT = Font(bold=True)
ERROR_FILL = PatternFill(fill_type="solid", fgColor="FCE4D6")


def copy_template_workbook(output_path: Path) -> Path:
    """Copy the canonical workbook template to a destination."""

    if not TEMPLATE_WORKBOOK_PATH.exists():
        raise FileNotFoundError(
            f"Workbook template not found at {TEMPLATE_WORKBOOK_PATH}. "
            "Run scripts/build_workbook_template.mjs to regenerate it."
        )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(TEMPLATE_WORKBOOK_PATH, output_path)
    return output_path


def load_workbook_project(path: Path) -> WorkbookProjectV1:
    """Load a workbook project into a normalized Python model."""

    workbook = load_workbook(path, data_only=True)
    missing = [sheet for sheet in ALL_WORKBOOK_SHEETS if sheet not in workbook.sheetnames]
    if missing:
        raise ValueError(f"Workbook is missing required sheets: {', '.join(missing)}")

    project_values = _parse_project_sheet(workbook[PROJECT_SHEET])
    settings = _parse_settings_sheet(workbook[SETTINGS_SHEET])
    staff = _parse_simple_table(
        workbook[STAFF_SHEET],
        expected_headers=["name", "roles", "target_hours", "max_hours", "year"],
        table_name="staff",
    )
    departments = _parse_simple_table(
        workbook[DEPARTMENTS_SHEET],
        expected_headers=["department", "target_hours", "max_hours"],
        table_name="departments",
    )
    unavailability = _parse_simple_table(
        workbook[UNAVAILABILITY_SHEET],
        expected_headers=["employee", "day", "start_time", "end_time", "buffer_before_start", "buffer_after_end"],
        table_name="unavailability",
    )

    preference_sections = {
        section: _parse_section_table(workbook[PREFERENCES_SHEET], section, headers)
        for section, headers in PREFERENCE_SECTION_HEADERS.items()
    }
    constraint_sections = {
        section: _parse_section_table(workbook["Constraints"], section, headers)
        for section, headers in CONSTRAINT_SECTION_HEADERS.items()
    }

    favored_employees = {
        _string(row["employee"]): _coerce_float(row["multiplier"], default=1.0)
        for row in preference_sections["Favored Employees"]
        if _string(row["employee"])
    }
    favored_departments = {
        normalize_department_name(_string(row["department"])): _coerce_float(row["multiplier"], default=1.0)
        for row in preference_sections["Favored Departments"]
        if _string(row["department"])
    }
    favored_frontdesk_departments = {
        normalize_department_name(_string(row["department"])): _coerce_float(row["multiplier"], default=1.0)
        for row in preference_sections["Favored Front Desk Departments"]
        if _string(row["department"])
    }
    favored_employee_departments = [
        FavoredEmployeeDepartment(
            employee=_string(row["employee"]),
            department=normalize_department_name(_string(row["department"])),
            multiplier=_coerce_float(row["multiplier"], default=1.0),
        )
        for row in preference_sections["Favored Employee Departments"]
        if _string(row["employee"]) or _string(row["department"])
    ]
    shift_time_preferences = [
        ShiftTimePreference(
            employee=_string(row["employee"]),
            day=_string(row["day"]),
            preference=_string(row["preference"]).lower(),
        )
        for row in preference_sections["Shift Time Preferences"]
        if _string(row["employee"]) or _string(row["day"]) or _string(row["preference"])
    ]
    training_requests = [
        TrainingRequest(
            department=normalize_department_name(_string(row["department"])),
            trainee_one=_string(row["trainee1"]),
            trainee_two=_string(row["trainee2"]),
        )
        for row in constraint_sections["Training Pairs"]
        if _string(row["department"]) or _string(row["trainee1"]) or _string(row["trainee2"])
    ]
    timeset_requests = [
        _build_timeset_request(row)
        for row in constraint_sections["Timesets"]
        if _string(row["employee"]) or _string(row["day"]) or _string(row["department"]) or _string(row["start_time"]) or _string(row["end_time"])
    ]
    equality_requests = [
        EqualityRequest(
            department=normalize_department_name(_string(row["department"])),
            employee1=_string(row["employee1"]),
            employee2=_string(row["employee2"]),
        )
        for row in constraint_sections["Equality Constraints"]
        if _string(row["department"]) or _string(row["employee1"]) or _string(row["employee2"])
    ]

    return WorkbookProjectV1(
        schema_version=int(project_values["schema_version"]),
        project_name=_string(project_values["project_name"]) or "Semester Scheduler Project",
        front_desk_enabled=_coerce_bool(project_values["front_desk_enabled"], default=True),
        settings=settings,
        staff=staff,
        departments=departments,
        unavailability=unavailability,
        favored_employees=favored_employees,
        favored_departments=favored_departments,
        favored_frontdesk_departments=favored_frontdesk_departments,
        favored_employee_departments=favored_employee_departments,
        shift_time_preferences=shift_time_preferences,
        training_requests=training_requests,
        timeset_requests=timeset_requests,
        equality_requests=equality_requests,
    )


def validate_workbook_project(project: WorkbookProjectV1) -> list[WorkbookValidationIssue]:
    """Validate workbook data and return structured issues."""

    issues: list[WorkbookValidationIssue] = []

    if project.schema_version != WORKBOOK_SCHEMA_VERSION:
        issues.append(
            WorkbookValidationIssue(
                severity="error",
                sheet=PROJECT_SHEET,
                table="project",
                row=None,
                column="schema_version",
                message=(
                    f"Unsupported workbook schema version {project.schema_version}. "
                    f"Expected {WORKBOOK_SCHEMA_VERSION}."
                ),
            )
        )

    staff_names: list[str] = []
    departments: list[str] = []

    if not project.staff:
        issues.append(_issue(STAFF_SHEET, "staff", None, "name", "At least one staff row is required."))
    if not project.departments:
        issues.append(_issue(DEPARTMENTS_SHEET, "departments", None, "department", "At least one department row is required."))

    for index, row in enumerate(project.staff, start=2):
        name = _string(row.get("name"))
        roles = _parse_roles(row.get("roles"))
        target_hours = _maybe_float(row.get("target_hours"))
        max_hours = _maybe_float(row.get("max_hours"))
        year = _maybe_float(row.get("year"))

        if not name:
            issues.append(_issue(STAFF_SHEET, "staff", index, "name", "Employee name is required."))
        else:
            staff_names.append(name)
        if not roles:
            issues.append(_issue(STAFF_SHEET, "staff", index, "roles", "At least one role is required."))
        if target_hours is None:
            issues.append(_issue(STAFF_SHEET, "staff", index, "target_hours", "Target hours must be numeric."))
        elif not is_slot_aligned_hours(target_hours):
            issues.append(_issue(STAFF_SHEET, "staff", index, "target_hours", "Target hours must align to the 10-minute grid."))
        if max_hours is None:
            issues.append(_issue(STAFF_SHEET, "staff", index, "max_hours", "Max hours must be numeric."))
        elif not is_slot_aligned_hours(max_hours):
            issues.append(_issue(STAFF_SHEET, "staff", index, "max_hours", "Max hours must align to the 10-minute grid."))
        if target_hours is not None and max_hours is not None and target_hours > max_hours:
            issues.append(_issue(STAFF_SHEET, "staff", index, "target_hours", "Target hours cannot exceed max hours."))
        if year is None or not float(year).is_integer():
            issues.append(_issue(STAFF_SHEET, "staff", index, "year", "Year must be an integer."))

    for normalized_name, count in Counter(name.strip().lower() for name in staff_names if name.strip()).items():
        if count > 1:
            issues.append(_issue(STAFF_SHEET, "staff", None, "name", f"Duplicate employee name detected after normalization: {normalized_name}."))

    for index, row in enumerate(project.departments, start=2):
        department = _string(row.get("department"))
        normalized = normalize_department_name(department)
        target_hours = _maybe_float(row.get("target_hours"))
        max_hours = _maybe_float(row.get("max_hours"))

        if not department:
            issues.append(_issue(DEPARTMENTS_SHEET, "departments", index, "department", "Department name is required."))
        else:
            departments.append(normalized)
        if target_hours is None:
            issues.append(_issue(DEPARTMENTS_SHEET, "departments", index, "target_hours", "Target hours must be numeric."))
        elif not is_slot_aligned_hours(target_hours):
            issues.append(_issue(DEPARTMENTS_SHEET, "departments", index, "target_hours", "Target hours must align to the 10-minute grid."))
        if max_hours is None:
            issues.append(_issue(DEPARTMENTS_SHEET, "departments", index, "max_hours", "Max hours must be numeric."))
        elif not is_slot_aligned_hours(max_hours):
            issues.append(_issue(DEPARTMENTS_SHEET, "departments", index, "max_hours", "Max hours must align to the 10-minute grid."))
        if target_hours is not None and max_hours is not None and target_hours > max_hours:
            issues.append(_issue(DEPARTMENTS_SHEET, "departments", index, "target_hours", "Target hours cannot exceed max hours."))

    for normalized_name, count in Counter(departments).items():
        if normalized_name and count > 1:
            issues.append(_issue(DEPARTMENTS_SHEET, "departments", None, "department", f"Duplicate department detected after normalization: {normalized_name}."))

    department_set = set(departments)
    staff_set = {name.strip().lower() for name in staff_names if name.strip()}

    for index, row in enumerate(project.staff, start=2):
        for role in _parse_roles(row.get("roles")):
            if role != "front_desk" and role not in department_set:
                issues.append(_issue(STAFF_SHEET, "staff", index, "roles", f"Unknown department role '{role}'."))

    for index, row in enumerate(project.unavailability, start=2):
        employee = _string(row.get("employee"))
        day = _string(row.get("day"))
        start_time = _string(row.get("start_time"))
        end_time = _string(row.get("end_time"))
        if employee and employee.strip().lower() not in staff_set:
            issues.append(_issue(UNAVAILABILITY_SHEET, "unavailability", index, "employee", f"Unknown employee '{employee}'."))
        if day and day not in DAY_NAMES:
            issues.append(_issue(UNAVAILABILITY_SHEET, "unavailability", index, "day", f"Day must be one of: {', '.join(DAY_NAMES)}."))
        if start_time and start_time not in TIME_SLOT_STARTS:
            issues.append(_issue(UNAVAILABILITY_SHEET, "unavailability", index, "start_time", "Start time must be on the 10-minute grid."))
        if end_time and end_time not in TIME_VALUE_OPTIONS:
            issues.append(_issue(UNAVAILABILITY_SHEET, "unavailability", index, "end_time", "End time must be on the 10-minute grid or the day end."))
        if start_time in TIME_SLOT_STARTS and end_time in TIME_VALUE_OPTIONS:
            if _time_to_slot_index(end_time, is_end=True) <= _time_to_slot_index(start_time):
                issues.append(_issue(UNAVAILABILITY_SHEET, "unavailability", index, "end_time", "End time must be after start time."))

    for employee, multiplier in project.favored_employees.items():
        if employee.strip().lower() not in staff_set:
            issues.append(_issue(PREFERENCES_SHEET, "Favored Employees", None, "employee", f"Unknown employee '{employee}'."))
        if not _is_finite_number(multiplier):
            issues.append(_issue(PREFERENCES_SHEET, "Favored Employees", None, "multiplier", f"Multiplier for '{employee}' must be numeric."))

    for mapping_name, values in (
        ("Favored Departments", project.favored_departments),
        ("Favored Front Desk Departments", project.favored_frontdesk_departments),
    ):
        for department, multiplier in values.items():
            if department not in department_set:
                issues.append(_issue(PREFERENCES_SHEET, mapping_name, None, "department", f"Unknown department '{department}'."))
            if not _is_finite_number(multiplier):
                issues.append(_issue(PREFERENCES_SHEET, mapping_name, None, "multiplier", f"Multiplier for '{department}' must be numeric."))

    for item in project.favored_employee_departments:
        if item.employee.strip().lower() not in staff_set:
            issues.append(_issue(PREFERENCES_SHEET, "Favored Employee Departments", None, "employee", f"Unknown employee '{item.employee}'."))
        if item.department != "front_desk" and item.department not in department_set:
            issues.append(_issue(PREFERENCES_SHEET, "Favored Employee Departments", None, "department", f"Unknown department '{item.department}'."))

    for item in project.shift_time_preferences:
        if item.employee.strip().lower() not in staff_set:
            issues.append(_issue(PREFERENCES_SHEET, "Shift Time Preferences", None, "employee", f"Unknown employee '{item.employee}'."))
        if item.day not in DAY_NAMES:
            issues.append(_issue(PREFERENCES_SHEET, "Shift Time Preferences", None, "day", f"Invalid day '{item.day}'."))
        if item.preference not in {"morning", "afternoon"}:
            issues.append(_issue(PREFERENCES_SHEET, "Shift Time Preferences", None, "preference", "Preference must be morning or afternoon."))

    for item in project.training_requests:
        if item.department not in department_set:
            issues.append(_issue("Constraints", "Training Pairs", None, "department", f"Unknown department '{item.department}'."))
        if item.trainee_one.strip().lower() not in staff_set:
            issues.append(_issue("Constraints", "Training Pairs", None, "trainee1", f"Unknown employee '{item.trainee_one}'."))
        if item.trainee_two.strip().lower() not in staff_set:
            issues.append(_issue("Constraints", "Training Pairs", None, "trainee2", f"Unknown employee '{item.trainee_two}'."))
        if item.trainee_one.strip().lower() == item.trainee_two.strip().lower():
            issues.append(_issue("Constraints", "Training Pairs", None, "trainee2", "Training pair employees must be different."))

    for item in project.timeset_requests:
        if item.employee.strip().lower() not in staff_set:
            issues.append(_issue("Constraints", "Timesets", None, "employee", f"Unknown employee '{item.employee}'."))
        if item.day not in DAY_NAMES:
            issues.append(_issue("Constraints", "Timesets", None, "day", f"Invalid day '{item.day}'."))
        if item.department != "front_desk" and item.department not in department_set:
            issues.append(_issue("Constraints", "Timesets", None, "department", f"Unknown department '{item.department}'."))
        if item.end_slot <= item.start_slot:
            issues.append(_issue("Constraints", "Timesets", None, "end_time", "Timeset end must be after start."))

    for item in project.equality_requests:
        if item.department not in department_set:
            issues.append(_issue("Constraints", "Equality Constraints", None, "department", f"Unknown department '{item.department}'."))
        if item.employee1.strip().lower() not in staff_set:
            issues.append(_issue("Constraints", "Equality Constraints", None, "employee1", f"Unknown employee '{item.employee1}'."))
        if item.employee2.strip().lower() not in staff_set:
            issues.append(_issue("Constraints", "Equality Constraints", None, "employee2", f"Unknown employee '{item.employee2}'."))
        if item.employee1.strip().lower() == item.employee2.strip().lower():
            issues.append(_issue("Constraints", "Equality Constraints", None, "employee2", "Equality pair employees must be different."))

    for definition in WORKBOOK_SETTINGS:
        value = project.settings.get(definition.key, definition.default)
        if definition.value_type == "int" and not _is_intlike(value):
            issues.append(_issue(SETTINGS_SHEET, "settings", None, definition.key, "Setting must be an integer."))
        elif definition.value_type == "float" and not _is_finite_number(value):
            issues.append(_issue(SETTINGS_SHEET, "settings", None, definition.key, "Setting must be numeric."))
        elif definition.value_type == "bool" and not isinstance(value, bool):
            issues.append(_issue(SETTINGS_SHEET, "settings", None, definition.key, "Setting must be TRUE or FALSE."))
        elif definition.value_type == "enum" and str(value) not in definition.choices:
            issues.append(
                _issue(
                    SETTINGS_SHEET,
                    "settings",
                    None,
                    definition.key,
                    f"Setting must be one of: {', '.join(definition.choices)}.",
                )
            )

    return issues


def solver_kwargs_from_project(project: WorkbookProjectV1) -> dict[str, Any]:
    """Translate workbook settings into solve_schedule keyword arguments."""

    kwargs: dict[str, Any] = {
        "solver_max_time": int(project.settings["solverMaxTime"]),
        "favored_employees": project.favored_employees,
        "training_requests": project.training_requests,
        "favored_departments": project.favored_departments,
        "favored_frontdesk_departments": project.favored_frontdesk_departments,
        "timeset_requests": project.timeset_requests,
        "favored_employee_depts": project.favored_employee_departments,
        "shift_time_preferences": project.shift_time_preferences,
        "equality_requests": project.equality_requests,
        "front_desk_enabled": project.front_desk_enabled,
        "show_progress": False,
    }

    for definition in WORKBOOK_SETTINGS:
        if definition.solver_argument is None or definition.key == "solverMaxTime":
            continue
        kwargs[definition.solver_argument] = project.settings.get(definition.key, definition.default)
    return kwargs


def materialize_solver_inputs(project: WorkbookProjectV1, target_dir: Path) -> tuple[Path, Path]:
    """Write temporary CSVs from workbook data for the existing solver path."""

    target_dir.mkdir(parents=True, exist_ok=True)
    staff_path = target_dir / "workbook_staff.csv"
    departments_path = target_dir / "workbook_departments.csv"

    unavailability_by_employee: dict[str, dict[str, list[dict[str, Any]]]] = defaultdict(lambda: {day: [] for day in DAY_NAMES})
    for row in project.unavailability:
        employee = _string(row.get("employee"))
        if not employee:
            continue
        day = _string(row.get("day"))
        if day not in DAY_NAMES:
            continue
        unavailability_by_employee[employee][day].append(
            {
                "startTime": _string(row.get("start_time")),
                "endTime": _string(row.get("end_time")),
                "bufferBeforeStart": _coerce_bool(row.get("buffer_before_start"), default=False),
                "bufferAfterEnd": _coerce_bool(row.get("buffer_after_end"), default=False),
            }
        )

    staff_rows: list[dict[str, Any]] = []
    for row in project.staff:
        name = _string(row.get("name"))
        if not name:
            continue
        row_payload: dict[str, Any] = {
            "name": name,
            "roles": ";".join(_parse_roles(row.get("roles"))),
            "target_hours": _maybe_float(row.get("target_hours")),
            "max_hours": _maybe_float(row.get("max_hours")),
            "year": int(_maybe_float(row.get("year")) or 0),
        }
        availability_by_day = _build_availability_grid(unavailability_by_employee.get(name, {day: [] for day in DAY_NAMES}))
        for day in DAY_NAMES:
            for slot_time, is_available in zip(TIME_SLOT_STARTS, availability_by_day[day]):
                row_payload[f"{day}_{slot_time}"] = 1 if is_available else 0
        staff_rows.append(row_payload)

    department_rows = [
        {
            "department": _string(row.get("department")),
            "target_hours": _maybe_float(row.get("target_hours")),
            "max_hours": _maybe_float(row.get("max_hours")),
        }
        for row in project.departments
        if _string(row.get("department"))
    ]

    pd.DataFrame(staff_rows).to_csv(staff_path, index=False)
    pd.DataFrame(department_rows).to_csv(departments_path, index=False)
    return staff_path, departments_path


def _build_availability_grid(day_blocks: dict[str, list[dict[str, Any]]]) -> dict[str, list[bool]]:
    availability_by_day = {day: [True] * len(TIME_SLOT_STARTS) for day in DAY_NAMES}
    for day in DAY_NAMES:
        for block in day_blocks.get(day, []):
            start = _string(block.get("startTime"))
            end = _string(block.get("endTime"))
            if start not in TIME_SLOT_STARTS or end not in TIME_VALUE_OPTIONS:
                continue
            start_slot = _time_to_slot_index(start)
            end_slot = _time_to_slot_index(end, is_end=True)
            for slot in range(start_slot, min(end_slot, len(TIME_SLOT_STARTS))):
                availability_by_day[day][slot] = False
            if _coerce_bool(block.get("bufferBeforeStart"), default=False) and start_slot > 0:
                availability_by_day[day][start_slot - 1] = False
            if _coerce_bool(block.get("bufferAfterEnd"), default=False) and end_slot < len(TIME_SLOT_STARTS):
                availability_by_day[day][end_slot] = False
    return availability_by_day


def export_solved_workbook(
    template_input_path: Path,
    output_path: Path,
    project: WorkbookProjectV1,
    validation_issues: list[WorkbookValidationIssue],
    status: int,
    solver_output_path: Path | None,
) -> Path:
    """Create the solved workbook copy with regenerated result sheets."""

    output_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_input_path, output_path)
    workbook = load_workbook(output_path)

    for sheet_name in [VALIDATION_SHEET, SOLVER_AVAILABILITY_SHEET, *RESULT_SHEET_NAMES.values()]:
        if sheet_name in workbook.sheetnames:
            del workbook[sheet_name]

    validation_sheet = workbook.create_sheet(VALIDATION_SHEET)
    _write_validation_sheet(validation_sheet, validation_issues)

    availability_sheet = workbook.create_sheet(SOLVER_AVAILABILITY_SHEET)
    _write_solver_availability_sheet(availability_sheet, project)
    availability_sheet.sheet_state = "hidden"

    result_sheet_order = [
        RESULT_SHEET_NAMES["status"],
        RESULT_SHEET_NAMES["formatted"],
        RESULT_SHEET_NAMES["weekly"],
        RESULT_SHEET_NAMES["monday"],
        RESULT_SHEET_NAMES["tuesday"],
        RESULT_SHEET_NAMES["wednesday"],
        RESULT_SHEET_NAMES["thursday"],
        RESULT_SHEET_NAMES["friday"],
        RESULT_SHEET_NAMES["employees"],
        RESULT_SHEET_NAMES["departments"],
    ]
    result_sheets = {name: workbook.create_sheet(name) for name in result_sheet_order}

    source_workbook = load_workbook(solver_output_path, data_only=True) if solver_output_path and solver_output_path.exists() else None
    formatted_output_path = (
        solver_output_path.with_name(f"{solver_output_path.stem}-formatted{solver_output_path.suffix}")
        if solver_output_path is not None
        else None
    )
    formatted_workbook = (
        load_workbook(formatted_output_path)
        if formatted_output_path is not None and formatted_output_path.exists()
        else None
    )
    _write_status_sheet(result_sheets[RESULT_SHEET_NAMES["status"]], project, status, source_workbook)

    source_mapping = {
        "Weekly Grid": RESULT_SHEET_NAMES["weekly"],
        "Mon Schedule": RESULT_SHEET_NAMES["monday"],
        "Tue Schedule": RESULT_SHEET_NAMES["tuesday"],
        "Wed Schedule": RESULT_SHEET_NAMES["wednesday"],
        "Thu Schedule": RESULT_SHEET_NAMES["thursday"],
        "Fri Schedule": RESULT_SHEET_NAMES["friday"],
        "Employee Summary": RESULT_SHEET_NAMES["employees"],
        "Department Summary": RESULT_SHEET_NAMES["departments"],
    }
    if formatted_workbook is not None and "Schedule" in formatted_workbook.sheetnames:
        _copy_sheet_with_formatting(formatted_workbook["Schedule"], result_sheets[RESULT_SHEET_NAMES["formatted"]])
    else:
        _write_placeholder_result_sheet(
            result_sheets[RESULT_SHEET_NAMES["formatted"]],
            "No formatted schedule workbook was produced.",
        )
    if source_workbook is not None:
        for source_name, target_name in source_mapping.items():
            if source_name in source_workbook.sheetnames:
                _copy_sheet_values(source_workbook[source_name], result_sheets[target_name], generated_label=source_name)
            else:
                _write_placeholder_result_sheet(result_sheets[target_name], f"Source sheet '{source_name}' was not produced.")
    else:
        for target_name in source_mapping.values():
            _write_placeholder_result_sheet(result_sheets[target_name], "No feasible solution was produced.")

    workbook.save(output_path)
    return output_path


def _parse_project_sheet(sheet: Worksheet) -> dict[str, Any]:
    values: dict[str, Any] = {}
    for row_idx in range(1, sheet.max_row + 1):
        key = _string(sheet.cell(row=row_idx, column=1).value)
        if not key:
            continue
        values[key] = sheet.cell(row=row_idx, column=2).value
    missing = [field for field in PROJECT_FIELDS if field not in values]
    if missing:
        raise ValueError(f"Project sheet is missing required keys: {', '.join(missing)}")
    return values


def _parse_settings_sheet(sheet: Worksheet) -> dict[str, Any]:
    rows = _parse_simple_table(
        sheet,
        expected_headers=["key", "default_value", "value", "unit", "description"],
        table_name="settings",
    )
    settings = {definition.key: definition.default for definition in WORKBOOK_SETTINGS}
    for row in rows:
        key = _string(row.get("key"))
        if not key or key not in WORKBOOK_SETTINGS_BY_KEY:
            continue
        definition = WORKBOOK_SETTINGS_BY_KEY[key]
        raw_value = row.get("value")
        settings[key] = _coerce_setting_value(raw_value, definition.default, definition.value_type)
    return settings


def _parse_simple_table(sheet: Worksheet, expected_headers: list[str], table_name: str) -> list[dict[str, Any]]:
    actual_headers = [_string(sheet.cell(row=1, column=col).value) for col in range(1, len(expected_headers) + 1)]
    if actual_headers != expected_headers:
        raise ValueError(
            f"{sheet.title} sheet headers do not match {table_name}: "
            f"expected {expected_headers}, got {actual_headers}"
        )
    rows: list[dict[str, Any]] = []
    for row_idx in range(2, sheet.max_row + 1):
        values = [sheet.cell(row=row_idx, column=col).value for col in range(1, len(expected_headers) + 1)]
        if not any(_string(value) for value in values):
            continue
        rows.append(dict(zip(expected_headers, values)))
    return rows


def _parse_section_table(sheet: Worksheet, title: str, headers: list[str]) -> list[dict[str, Any]]:
    for row_idx in range(1, sheet.max_row + 1):
        if _string(sheet.cell(row=row_idx, column=1).value) != title:
            continue
        header_row = row_idx + 1
        actual_headers = [_string(sheet.cell(row=header_row, column=col).value) for col in range(1, len(headers) + 1)]
        if actual_headers != headers:
            raise ValueError(
                f"{sheet.title} section '{title}' headers do not match: expected {headers}, got {actual_headers}"
            )
        rows: list[dict[str, Any]] = []
        data_row = header_row + 1
        while data_row <= sheet.max_row:
            values = [sheet.cell(row=data_row, column=col).value for col in range(1, len(headers) + 1)]
            if not any(_string(value) for value in values):
                break
            rows.append(dict(zip(headers, values)))
            data_row += 1
        return rows
    raise ValueError(f"Missing expected section '{title}' in {sheet.title}.")


def _build_timeset_request(row: dict[str, Any]) -> TimesetRequest:
    start_label = _string(row["start_time"])
    end_label = _string(row["end_time"])
    return TimesetRequest(
        employee=_string(row["employee"]),
        day=_string(row["day"]),
        department=normalize_department_name(_string(row["department"])),
        start_slot=_time_to_slot_index(start_label),
        end_slot=_time_to_slot_index(end_label, is_end=True),
    )


def _time_to_slot_index(label: str, *, is_end: bool = False) -> int:
    if is_end and label == FINAL_EDGE_LABEL:
        return len(TIME_SLOT_STARTS)
    return TIME_SLOT_STARTS.index(label)


def _coerce_setting_value(value: Any, default: Any, value_type: str) -> Any:
    if value is None or _string(value) == "":
        return default
    if value_type == "bool":
        return _coerce_bool(value, default=bool(default))
    if value_type == "int":
        parsed = _maybe_float(value)
        return int(parsed) if parsed is not None else default
    if value_type == "float":
        parsed = _maybe_float(value)
        return float(parsed) if parsed is not None else default
    if value_type == "enum":
        return _string(value) or default
    return value


def _write_validation_sheet(sheet: Worksheet, issues: list[WorkbookValidationIssue]) -> None:
    sheet["A1"] = "Workbook Validation"
    sheet["A1"].font = Font(size=16, bold=True)
    error_count = sum(1 for issue in issues if issue.severity == "error")
    warning_count = sum(1 for issue in issues if issue.severity == "warning")
    sheet["A3"] = "Status"
    sheet["B3"] = "PASS" if error_count == 0 else "FAILED"
    sheet["A4"] = "Errors"
    sheet["B4"] = error_count
    sheet["A5"] = "Warnings"
    sheet["B5"] = warning_count

    headers = ["Severity", "Sheet", "Table", "Row", "Column", "Message"]
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=7, column=index, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    for row_index, issue in enumerate(issues, start=8):
        values = [issue.severity, issue.sheet, issue.table, issue.row, issue.column, issue.message]
        for col_index, value in enumerate(values, start=1):
            cell = sheet.cell(row=row_index, column=col_index, value=value)
            if issue.severity == "error":
                cell.fill = ERROR_FILL
    sheet.freeze_panes = "A7"


def _write_solver_availability_sheet(sheet: Worksheet, project: WorkbookProjectV1) -> None:
    headers = ["employee", "day", *TIME_SLOT_STARTS]
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL

    blocks_by_employee_day: dict[tuple[str, str], list[dict[str, Any]]] = defaultdict(list)
    for row in project.unavailability:
        employee = _string(row.get("employee"))
        day = _string(row.get("day"))
        if not employee or day not in DAY_NAMES:
            continue
        blocks_by_employee_day[(employee, day)].append(row)

    row_index = 2
    for staff_row in project.staff:
        employee = _string(staff_row.get("name"))
        if not employee:
            continue
        for day in DAY_NAMES:
            sheet.cell(row=row_index, column=1, value=employee)
            sheet.cell(row=row_index, column=2, value=day)
            available = [True] * len(TIME_SLOT_STARTS)
            for block in blocks_by_employee_day.get((employee, day), []):
                start = _string(block.get("start_time"))
                end = _string(block.get("end_time"))
                if start not in TIME_SLOT_STARTS or end not in TIME_VALUE_OPTIONS:
                    continue
                start_slot = _time_to_slot_index(start)
                end_slot = _time_to_slot_index(end, is_end=True)
                for slot in range(start_slot, min(end_slot, len(available))):
                    available[slot] = False
                if _coerce_bool(block.get("buffer_before_start"), default=False) and start_slot > 0:
                    available[start_slot - 1] = False
                if _coerce_bool(block.get("buffer_after_end"), default=False) and end_slot < len(available):
                    available[end_slot] = False
            for col_index, is_available in enumerate(available, start=3):
                sheet.cell(row=row_index, column=col_index, value=1 if is_available else 0)
            row_index += 1
    sheet.freeze_panes = "C2"


def _write_status_sheet(
    sheet: Worksheet,
    project: WorkbookProjectV1,
    status: int,
    source_workbook: Workbook | None,
) -> None:
    status_label = {
        cp_model.OPTIMAL: "OPTIMAL",
        cp_model.FEASIBLE: "FEASIBLE",
        cp_model.INFEASIBLE: "INFEASIBLE",
        cp_model.MODEL_INVALID: "MODEL_INVALID",
        cp_model.UNKNOWN: "UNKNOWN",
    }.get(status, str(status))
    sheet["A1"] = "Solve Status"
    sheet["A1"].font = Font(size=16, bold=True)
    rows = [
        ("Project Name", project.project_name),
        ("Schema Version", project.schema_version),
        ("Front Desk Enabled", project.front_desk_enabled),
        ("Solver Status", status_label),
        ("Solver Max Seconds", project.settings.get("solverMaxTime")),
    ]
    for row_index, (label, value) in enumerate(rows, start=3):
        sheet.cell(row=row_index, column=1, value=label).font = HEADER_FONT
        sheet.cell(row=row_index, column=2, value=value)

    if source_workbook is not None and "Role Distribution" in source_workbook.sheetnames:
        sheet["A8"] = "Role Distribution"
        sheet["A8"].font = HEADER_FONT
        _append_sheet_values(source_workbook["Role Distribution"], sheet, start_row=9)
    else:
        sheet["A8"] = "No role distribution table was available."
    sheet.freeze_panes = "A3"


def _copy_sheet_values(source: Worksheet, target: Worksheet, generated_label: str) -> None:
    target["A1"] = f"Generated: {generated_label}"
    target["A1"].font = HEADER_FONT
    _append_sheet_values(source, target, start_row=3)
    target.freeze_panes = "A3"


def _copy_sheet_with_formatting(source: Worksheet, target: Worksheet) -> None:
    target.sheet_view.showGridLines = source.sheet_view.showGridLines
    target.freeze_panes = source.freeze_panes

    for row in source.iter_rows():
        for cell in row:
            target_cell = target[cell.coordinate]
            target_cell.value = cell.value
            if cell.has_style:
                target_cell._style = copy(cell._style)
            if cell.number_format:
                target_cell.number_format = cell.number_format
            if cell.font:
                target_cell.font = copy(cell.font)
            if cell.fill:
                target_cell.fill = copy(cell.fill)
            if cell.border:
                target_cell.border = copy(cell.border)
            if cell.alignment:
                target_cell.alignment = copy(cell.alignment)
            if cell.protection:
                target_cell.protection = copy(cell.protection)
            if cell.comment is not None:
                target_cell.comment = copy(cell.comment)
            if cell.hyperlink:
                target_cell._hyperlink = copy(cell.hyperlink)

    for merged_range in source.merged_cells.ranges:
        target.merge_cells(str(merged_range))

    for key, dimension in source.column_dimensions.items():
        target_dimension = target.column_dimensions[key]
        target_dimension.width = dimension.width
        target_dimension.hidden = dimension.hidden
        target_dimension.bestFit = dimension.bestFit
        target_dimension.outlineLevel = dimension.outlineLevel

    for key, dimension in source.row_dimensions.items():
        target_dimension = target.row_dimensions[key]
        target_dimension.height = dimension.height
        target_dimension.hidden = dimension.hidden
        target_dimension.outlineLevel = dimension.outlineLevel


def _append_sheet_values(source: Worksheet, target: Worksheet, start_row: int) -> None:
    for row_offset, row in enumerate(source.iter_rows(values_only=True), start=0):
        for col_index, value in enumerate(row, start=1):
            target.cell(row=start_row + row_offset, column=col_index, value=value)
            if row_offset == 0:
                target.cell(row=start_row + row_offset, column=col_index).font = HEADER_FONT
                target.cell(row=start_row + row_offset, column=col_index).fill = HEADER_FILL


def _write_placeholder_result_sheet(sheet: Worksheet, message: str) -> None:
    sheet["A1"] = "Generated Result"
    sheet["A1"].font = Font(size=16, bold=True)
    sheet["A3"] = message


def _parse_roles(raw_roles: Any) -> list[str]:
    text = _string(raw_roles)
    if not text:
        return []
    return [normalize_department_name(part) for part in text.replace(";", ",").split(",") if part.strip()]


def _string(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _maybe_float(value: Any) -> float | None:
    text = _string(value)
    if not text:
        return None
    try:
        number = float(text)
    except ValueError:
        return None
    if math.isnan(number) or math.isinf(number):
        return None
    return number


def _coerce_float(value: Any, *, default: float) -> float:
    parsed = _maybe_float(value)
    return default if parsed is None else parsed


def _coerce_bool(value: Any, *, default: bool) -> bool:
    if isinstance(value, bool):
        return value
    text = _string(value).lower()
    if text in {"true", "yes", "1", "y"}:
        return True
    if text in {"false", "no", "0", "n"}:
        return False
    return default


def _is_intlike(value: Any) -> bool:
    parsed = _maybe_float(value)
    return parsed is not None and float(parsed).is_integer()


def _is_finite_number(value: Any) -> bool:
    return _maybe_float(value) is not None


def _issue(sheet: str, table: str, row: int | None, column: str | None, message: str) -> WorkbookValidationIssue:
    return WorkbookValidationIssue(
        severity="error",
        sheet=sheet,
        table=table,
        row=row,
        column=column,
        message=message,
    )
